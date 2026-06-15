#!/usr/bin/env python3
"""
Golden-file regression tests for PDF → Word / Excel / PowerPoint engines.

Usage (from repo root):
  python Pdf-Engine/tests/run_golden_tests.py

Exit code 0 = all checks passed; 1 = regression detected.
Requires: pip install -r Pdf-Engine/requirements.txt
"""

from __future__ import annotations

import os
import re
import sys
import tempfile
import traceback
from dataclasses import dataclass

# Allow imports from Pdf-Engine root and tests/
_ENGINE = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
_TESTS = os.path.dirname(__file__)
for _p in (_ENGINE, _TESTS):
    if _p not in sys.path:
        sys.path.insert(0, _p)

from golden_fixtures import ensure_fixtures  # noqa: E402


@dataclass
class CaseResult:
    name: str
    passed: bool
    detail: str


def _docx_text(path: str) -> str:
    from docx import Document

    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def _xlsx_stats(path: str) -> tuple[int, int]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = len(wb.sheetnames)
    cells = 0
    for name in wb.sheetnames:
        ws = wb[name]
        for row in ws.iter_rows(values_only=True):
            for val in row:
                if val is not None and str(val).strip():
                    cells += 1
    wb.close()
    return sheets, cells


def _pptx_slides(path: str) -> int:
    from pptx import Presentation

    return len(Presentation(path).slides)


def _text_overlap_ratio(source: str, output: str) -> float:
    src_tokens = set(re.findall(r"[a-z0-9]+", source.lower()))
    out_tokens = set(re.findall(r"[a-z0-9]+", output.lower()))
    if not src_tokens:
        return 0.0
    return len(src_tokens & out_tokens) / len(src_tokens)


def _pdf_source_text(pdf_path: str) -> str:
    import fitz

    doc = fitz.open(pdf_path)
    try:
        return "\n".join(page.get_text("text") for page in doc)
    finally:
        doc.close()


def test_pdf_to_docx(pdf_path: str, work_dir: str) -> CaseResult:
    from pdf_to_docx_converter import convert

    out = os.path.join(work_dir, "out.docx")
    try:
        convert(pdf_path, out)
        if not os.path.isfile(out) or os.path.getsize(out) < 2000:
            return CaseResult("pdf_to_docx", False, "DOCX missing or too small")
        src = _pdf_source_text(pdf_path)
        docx = _docx_text(out)
        ratio = _text_overlap_ratio(src, docx)
        keywords = ("widget", "item", "shipping", "ratpdf", "golden")
        keyword_hit = any(k in docx.lower() for k in keywords)
        if ratio < 0.45 and not keyword_hit:
            return CaseResult("pdf_to_docx", False, f"text overlap {ratio:.0%} < 45%")
        return CaseResult("pdf_to_docx", True, f"text overlap {ratio:.0%}")
    except Exception as exc:
        return CaseResult("pdf_to_docx", False, str(exc))


def test_pdf_to_xlsx(pdf_path: str, work_dir: str) -> CaseResult:
    from office_converter import pdf_to_xlsx

    out = os.path.join(work_dir, "out.xlsx")
    try:
        pdf_to_xlsx(pdf_path, out)
        if not os.path.isfile(out) or os.path.getsize(out) < 1500:
            return CaseResult("pdf_to_xlsx", False, "XLSX missing or too small")
        sheets, cells = _xlsx_stats(out)
        if sheets < 1 or cells < 3:
            return CaseResult("pdf_to_xlsx", False, f"only {sheets} sheet(s), {cells} cell(s)")
        return CaseResult("pdf_to_xlsx", True, f"{sheets} sheet(s), {cells} non-empty cell(s)")
    except Exception as exc:
        return CaseResult("pdf_to_xlsx", False, str(exc))


def test_pdf_to_pptx(pdf_path: str, work_dir: str, expected_slides: int) -> CaseResult:
    from office_converter import pdf_to_pptx

    out = os.path.join(work_dir, "out.pptx")
    try:
        pdf_to_pptx(pdf_path, out)
        if not os.path.isfile(out) or os.path.getsize(out) < 5000:
            return CaseResult("pdf_to_pptx", False, "PPTX missing or too small")
        slides = _pptx_slides(out)
        if slides != expected_slides:
            return CaseResult("pdf_to_pptx", False, f"{slides} slides != {expected_slides} expected")
        return CaseResult("pdf_to_pptx", True, f"{slides} slide(s)")
    except Exception as exc:
        return CaseResult("pdf_to_pptx", False, str(exc))


def main() -> int:
    fixtures_dir = os.path.join(os.path.dirname(__file__), "golden", "fixtures")
    fixtures = ensure_fixtures(fixtures_dir)

    matrix = [
        ("simple_text", fixtures["simple_text"], 1, [test_pdf_to_docx, test_pdf_to_xlsx, test_pdf_to_pptx]),
        ("table_grid", fixtures["table_grid"], 1, [test_pdf_to_docx, test_pdf_to_xlsx, test_pdf_to_pptx]),
        ("two_page", fixtures["two_page"], 2, [test_pdf_to_pptx]),
    ]

    results: list[CaseResult] = []
    with tempfile.TemporaryDirectory(prefix="ratpdf_golden_") as work_root:
        for label, pdf_path, slide_count, tests in matrix:
            case_dir = os.path.join(work_root, label)
            os.makedirs(case_dir, exist_ok=True)
            for fn in tests:
                try:
                    if fn is test_pdf_to_pptx:
                        res = fn(pdf_path, case_dir, slide_count)
                        res = CaseResult(f"{label}/{res.name}", res.passed, res.detail)
                    else:
                        res = fn(pdf_path, case_dir)
                        res = CaseResult(f"{label}/{res.name}", res.passed, res.detail)
                except Exception:
                    res = CaseResult(f"{label}/{fn.__name__}", False, traceback.format_exc(limit=2))
                results.append(res)

    passed = sum(1 for r in results if r.passed)
    total = len(results)
    print(f"\nRatPDF golden tests: {passed}/{total} passed\n")
    for r in results:
        mark = "PASS" if r.passed else "FAIL"
        print(f"  [{mark}] {r.name}: {r.detail}")

    return 0 if passed == total else 1


if __name__ == "__main__":
    raise SystemExit(main())
