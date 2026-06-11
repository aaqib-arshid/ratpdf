"""
Build a DOCX from parsed layout, tables, and images using python-docx.

Uses section columns for multi-column pages and preserves spacing/reading order.
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import List, Optional, Tuple, Union

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from image_extractor import ExtractedImage, images_by_page
from pdf_parser import DocumentLayout, PageLayout, TextSpan, group_into_paragraphs, reading_order_key
from table_extractor import ExtractedTable, tables_by_page


# PDF points → inches
def pt_to_inches(pt: float) -> float:
    return pt / 72.0


def hex_to_rgb(hex_color: str) -> Optional[RGBColor]:
    if not hex_color or not hex_color.startswith("#") or len(hex_color) < 7:
        return None
    try:
        r = int(hex_color[1:3], 16)
        g = int(hex_color[3:5], 16)
        b = int(hex_color[5:7], 16)
        return RGBColor(r, g, b)
    except ValueError:
        return None


@dataclass
class PageElement:
    kind: str  # text | table | image
    y: float
    x: float
    column: int
    payload: Union[List[List[TextSpan]], ExtractedTable, ExtractedImage]


def _set_section_columns(section, num_cols: int, space_twips: int = 720) -> None:
    sect_pr = section._sectPr
    cols = sect_pr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(num_cols))
    cols.set(qn("w:space"), str(space_twips))


def _apply_run_style(run, span: TextSpan) -> None:
    run.font.name = span.font_family
    run.font.size = Pt(max(span.font_size, 6.0))
    run.bold = span.bold
    run.italic = span.italic
    color = hex_to_rgb(span.color)
    if color:
        run.font.color.rgb = color


def _add_text_paragraph(
    doc: Document,
    paragraph_spans: List[TextSpan],
    prev_bottom: Optional[float],
) -> float:
    """Add a paragraph from grouped spans; return new bottom Y."""
    if not paragraph_spans:
        return prev_bottom or 0.0

    para = doc.add_paragraph()
    top = min(s.y for s in paragraph_spans)
    if prev_bottom is not None:
        gap_pt = top - prev_bottom
        if gap_pt > 2:
            para.paragraph_format.space_before = Pt(min(gap_pt * 0.85, 72))

    avg_size = sum(s.font_size for s in paragraph_spans) / len(paragraph_spans)
    if avg_size >= 16:
        para.style = doc.styles["Heading 1"] if avg_size >= 20 else doc.styles["Heading 2"]

    ordered = sorted(paragraph_spans, key=lambda s: (s.x,))
    for i, span in enumerate(ordered):
        if i > 0:
            prev = ordered[i - 1]
            h_gap = span.x - (prev.x + prev.width)
            if h_gap > max(prev.font_size, span.font_size) * 0.4:
                para.add_run("\t")
        run = para.add_run(span.text)
        _apply_run_style(run, span)

    bottom = max(s.y + s.height for s in paragraph_spans)
    return bottom


def _add_table(doc: Document, table: ExtractedTable, prev_bottom: Optional[float]) -> float:
    if prev_bottom is not None:
        gap_pt = table.y - prev_bottom
        if gap_pt > 2:
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_before = Pt(min(gap_pt * 0.85, 72))

    row_count = len(table.rows)
    col_count = max(len(r) for r in table.rows) if table.rows else 0
    if row_count == 0 or col_count == 0:
        return table.y + table.height

    word_table = doc.add_table(rows=row_count, cols=col_count)
    word_table.style = "Table Grid"
    word_table.autofit = False

    for r_idx, row in enumerate(table.rows):
        for c_idx in range(col_count):
            text = row[c_idx] if c_idx < len(row) else ""
            cell = word_table.rows[r_idx].cells[c_idx]
            cell.text = text

            matching = [
                c for c in table.cells
                if c.row == r_idx and c.col == c_idx
            ]
            if matching and cell.paragraphs:
                para = cell.paragraphs[0]
                run = para.runs[0] if para.runs else para.add_run(text)
                mc = matching[0]
                run.font.name = mc.font_family
                run.font.size = Pt(max(mc.font_size, 6.0))
                run.bold = mc.bold

    return table.y + table.height


def _add_image(doc: Document, image: ExtractedImage, prev_bottom: Optional[float]) -> float:
    if prev_bottom is not None:
        gap_pt = image.y - prev_bottom
        if gap_pt > 2:
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_before = Pt(min(gap_pt * 0.85, 72))

    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = para.add_run()
    width_in = pt_to_inches(image.width)
    height_in = pt_to_inches(image.height)
    max_width = 6.5
    if width_in > max_width:
        scale = max_width / width_in
        width_in *= scale
        height_in *= scale
    run.add_picture(BytesIO(image.image_bytes), width=Inches(width_in), height=Inches(height_in))
    return image.y + image.height


def _collect_page_elements(
    page: PageLayout,
    page_tables: List[ExtractedTable],
    page_images: List[ExtractedImage],
) -> List[PageElement]:
    elements: List[PageElement] = []

    paragraphs = group_into_paragraphs(page.body_spans)
    for para_spans in paragraphs:
        if not para_spans:
            continue
        elements.append(
            PageElement(
                kind="text",
                y=para_spans[0].y,
                x=para_spans[0].x,
                column=para_spans[0].column,
                payload=para_spans,
            )
        )

    for table in page_tables:
        elements.append(
            PageElement(
                kind="table",
                y=table.y,
                x=table.x,
                column=0,
                payload=table,
            )
        )

    for image in page_images:
        elements.append(
            PageElement(
                kind="image",
                y=image.y,
                x=image.x,
                column=0,
                payload=image,
            )
        )

    elements.sort(key=lambda e: (e.column, e.y, e.x))
    return elements


def _add_header_footer(
    doc: Document,
    header_spans: List[TextSpan],
    footer_spans: List[TextSpan],
) -> None:
    section = doc.sections[-1]
    if header_spans:
        header = section.header
        hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        hp._element.clear()
        for span in sorted(header_spans, key=reading_order_key):
            run = hp.add_run(span.text + " ")
            _apply_run_style(run, span)
        hp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if footer_spans:
        footer = section.footer
        fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        fp._element.clear()
        for span in sorted(footer_spans, key=reading_order_key):
            run = fp.add_run(span.text + " ")
            _apply_run_style(run, span)
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER


def build(
    layout: DocumentLayout,
    tables: List[ExtractedTable],
    images: List[ExtractedImage],
    output_path: str,
) -> str:
    """Build DOCX from layout model and write to output_path."""
    doc = Document()
    tbl_map = tables_by_page(tables)
    img_map = images_by_page(images)

    for page_idx, page in enumerate(layout.pages):
        if page_idx > 0:
            pb = doc.add_paragraph()
            pb.add_run().add_break(WD_BREAK.PAGE)

        section = doc.sections[-1]
        num_cols = len(page.column_boundaries) if page.is_multi_column else 1
        if num_cols > 1:
            _set_section_columns(section, num_cols)

        _add_header_footer(doc, page.header_spans, page.footer_spans)

        page_tables = tbl_map.get(page.page_num, [])
        page_images = img_map.get(page.page_num, [])
        elements = _collect_page_elements(page, page_tables, page_images)

        prev_bottom: Optional[float] = None
        for element in elements:
            if element.kind == "text":
                prev_bottom = _add_text_paragraph(doc, element.payload, prev_bottom)
            elif element.kind == "table":
                prev_bottom = _add_table(doc, element.payload, prev_bottom)
            elif element.kind == "image":
                prev_bottom = _add_image(doc, element.payload, prev_bottom)

    doc.save(output_path)
    return output_path
