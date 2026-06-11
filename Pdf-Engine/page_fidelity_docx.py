"""
Visual-perfect PDF → DOCX: embed each page as a high-resolution image.
"""

from __future__ import annotations

import os
import zipfile
from io import BytesIO

import fitz
from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches


def count_pdf_images(pdf_path: str) -> int:
    doc = fitz.open(pdf_path)
    try:
        return sum(len(page.get_images(full=True)) for page in doc)
    finally:
        doc.close()


def count_docx_images(docx_path: str) -> int:
    if not os.path.isfile(docx_path):
        return 0
    try:
        with zipfile.ZipFile(docx_path) as zf:
            return sum(1 for n in zf.namelist() if n.startswith("word/media/"))
    except Exception:
        return 0


def docx_text_length(docx_path: str) -> int:
    try:
        from docx import Document as Doc
        doc = Doc(docx_path)
        parts: list[str] = []
        for para in doc.paragraphs:
            if para.text.strip():
                parts.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        parts.append(cell.text)
        return len("".join(parts).strip())
    except Exception:
        return 0


def docx_has_content(docx_path: str) -> bool:
    """True if DOCX has meaningful text OR embedded images (page renders count)."""
    if not os.path.isfile(docx_path):
        return False
    size = os.path.getsize(docx_path)
    if size < 2500:
        return False
    if count_docx_images(docx_path) > 0:
        return True
    return docx_text_length(docx_path) >= 10


def visual_fidelity_ok(pdf_path: str, docx_path: str) -> bool:
    pdf_imgs = count_pdf_images(pdf_path)
    if pdf_imgs == 0:
        return True
    docx_imgs = count_docx_images(docx_path)
    return docx_imgs >= max(1, int(pdf_imgs * 0.5))


def _render_page_pixmap(page: fitz.Page, dpi: int) -> bytes:
    """Render page to PNG bytes with fallbacks for masked/CMYK PDFs."""
    matrix = fitz.Matrix(dpi / 72.0, dpi / 72.0)
    attempts = [
        {"matrix": matrix, "alpha": False},
        {"matrix": matrix, "alpha": False, "colorspace": fitz.csRGB},
        {"dpi": dpi, "alpha": False},
    ]
    last_err: Exception | None = None
    for kwargs in attempts:
        try:
            pix = page.get_pixmap(**kwargs)
            return pix.tobytes("png")
        except Exception as exc:
            last_err = exc
            continue
    raise RuntimeError(f"Could not render page: {last_err}")


def build_pixmap_docx(pdf_path: str, output_docx_path: str, dpi: int = 200) -> str:
    """Render every PDF page at high DPI and embed in Word (pixel-perfect)."""
    pdf = fitz.open(pdf_path)
    word = Document()
    usable_width_in = 6.5

    try:
        for page_num, page in enumerate(pdf):
            if page_num > 0:
                pb = word.add_paragraph()
                pb.add_run().add_break(WD_BREAK.PAGE)

            img_bytes = _render_page_pixmap(page, dpi)
            page_w = page.rect.width or 1.0
            page_h = page.rect.height or 1.0
            height_in = usable_width_in * (page_h / page_w)

            para = word.add_paragraph()
            run = para.add_run()
            run.add_picture(
                BytesIO(img_bytes),
                width=Inches(usable_width_in),
                height=Inches(height_in),
            )
    finally:
        pdf.close()

    os.makedirs(os.path.dirname(os.path.abspath(output_docx_path)) or ".", exist_ok=True)
    word.save(output_docx_path)
    if not docx_has_content(output_docx_path):
        raise RuntimeError("Pixmap DOCX was created but contains no usable content")
    return output_docx_path
