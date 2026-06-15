#!/usr/bin/env python3
"""PDF → PPTX via full-slide page images (Adobe/iLovePDF-style fidelity for all PDF types)."""

from __future__ import annotations

import os

import fitz
from pptx import Presentation
from pptx.util import Emu


def _slide_size_emu(page_width_pt: float, page_height_pt: float) -> tuple[int, int]:
    """Match slide aspect ratio to the PDF page (10\" width baseline)."""
    width_in = 10.0
    height_in = width_in * (page_height_pt / max(page_width_pt, 1.0))
    return int(width_in * 914400), int(height_in * 914400)


def pdf_to_pptx_pixmap(pdf_path: str, output_path: str, dpi: int = 150) -> None:
    doc = fitz.open(pdf_path)
    try:
        if doc.page_count == 0:
            raise RuntimeError("PDF has no pages")

        first = doc[0]
        slide_w, slide_h = _slide_size_emu(first.rect.width, first.rect.height)

        prs = Presentation()
        prs.slide_width = Emu(slide_w)
        prs.slide_height = Emu(slide_h)
        blank = prs.slide_layouts[6]

        zoom = dpi / 72.0
        matrix = fitz.Matrix(zoom, zoom)
        work_dir = os.path.dirname(os.path.abspath(output_path)) or "."
        temp_images: list[str] = []

        for i, page in enumerate(doc):
            slide = prs.slides.add_slide(blank)
            pix = page.get_pixmap(matrix=matrix, alpha=False)
            img_path = os.path.join(work_dir, f".ratpdf_slide_{i:04d}.png")
            pix.save(img_path)
            pix = None
            temp_images.append(img_path)
            slide.shapes.add_picture(img_path, 0, 0, width=Emu(slide_w), height=Emu(slide_h))

        out_dir = os.path.dirname(os.path.abspath(output_path)) or "."
        os.makedirs(out_dir, exist_ok=True)
        prs.save(output_path)
        for img_path in temp_images:
            try:
                os.remove(img_path)
            except OSError:
                pass
    finally:
        doc.close()
